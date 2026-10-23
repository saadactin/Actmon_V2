"""The audit-ready compliance export: evidence, exceptions, integrity, formats.

No database. Every assertion goes through the production functions —
`compliance.build`, `compliance.control_evidence`, `assemble_pack`, the three
renderers and the route handler — never a re-implementation of them (see the
"test mirrors prove nothing" lesson). What this suite pins, and why:

* **The pack cannot disagree with the Compliance tab.** Every control state
  and every framework number in the pack is the one `compliance.build`
  resolved, and each control's violation evidence is exactly its
  `failing_findings`.
* **Exceptions are never hidden.** A HIGH finding accepted as a risk makes its
  control read PASS — and the pack flags that control as passed only with
  exceptions, lists the decision with who/when/why, subtracts it from
  "conforming", and flags an exception that has no recorded reason.
* **Severity bars hold.** A LOW finding fails a strict CIS control but is
  only advisory against ISO 27001.
* **Other clouds' checks are not evidence.** SOC 2 CC6.1 on an OCI account
  lists the OCI bucket check, not the S3 and Azure ones that cannot run.
* **Tamper-evident and reproducible.** The SHA-256 verifies, survives a JSON
  round trip, is deterministic, and breaks on any edit.
* **Formats.** The PDF carries the pack ID and the exception flags; the
  workbook has every control and evidence row, writes a formula-shaped
  resource name as text, and adds raw evidence only when asked; the ZIP's
  checksums match its members.
* **The route.** `/compliance/export` is never matched as a control named
  "export"; unknown formats/frameworks are 400s; a viewer (no
  reports:export) and a raw request without security:view_raw are audited
  403s.
"""
import asyncio
import base64
import hashlib
import io
import json
import pathlib
import re
import sys
import uuid
import zipfile
import zlib
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))

import logging  # noqa: E402

logging.disable(logging.CRITICAL)

from fastapi import HTTPException  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app.api import rbac  # noqa: E402
from app.security import compliance as comp  # noqa: E402
from app.security import controls as ctl  # noqa: E402
from app.security import registry as reg  # noqa: E402
from app.security.model import (  # noqa: E402
    ACCEPTED_RISK, CRITICAL, HIGH, LOW, MEDIUM, NOT_SCANNED, OPEN, SAFE, SUPPRESSED,
    UNKNOWN, VULNERABLE, SecurityFinding,
)
from app.services import compliance_export_service as cx  # noqa: E402

ok, fail = [], []


def check(label, cond, detail=""):
    (ok if cond else fail).append(label)
    print(f"  [{'PASS' if cond else 'FAIL'}] {label}" + (f"  {detail}" if detail else ""))


def eq(label, got, want):
    check(label, got == want, f"got {got!r}, want {want!r}")


NOW = datetime(2026, 10, 1, 9, 30, tzinfo=timezone.utc)
PROVIDER = "OCI"


def finding(rule_id, severity, *, name, verdict=VULNERABLE, status=OPEN, reason=None,
            by=None, affected=1):
    rule = reg.by_id(rule_id)
    return SecurityFinding(
        rule_id=rule_id, rule_version="1.0", provider=PROVIDER, account_id="a",
        fingerprint=uuid.uuid4().hex, resource_key=f"ocid1.test.{uuid.uuid4().hex[:10]}",
        resource_name=name, resource_type=rule.resource_types[0], region="ap-mumbai-1",
        title=rule.title, description=f"{rule.title} on {name}", remediation=rule.remediation,
        severity=severity, verdict=verdict, status=status, affected_count=affected,
        status_reason=reason, status_changed_by=by,
        status_changed_at=(NOW - timedelta(days=3)) if status != OPEN else None,
        first_detected_at=NOW - timedelta(days=20), last_detected_at=NOW - timedelta(hours=2),
        compliance=ctl.mappings_for_rule(rule_id, provider=PROVIDER))


# Rule statistics are an INPUT here (production derives them from inventory in
# SecurityCenterService._rule_stats); they are written in both shapes the
# compliance module reads — per rule, and per (rule, resource type).
def stats(rule_stats, rule_id, **counts):
    entry = {k: v for k, v in counts.items() if v}
    rule_stats[rule_id] = dict(entry)
    for rtype in reg.by_id(rule_id).resource_types:
        rule_stats[ctl.stat_key(rule_id, rtype)] = dict(entry)


F_SSH = finding("OCI_SECLIST_SSH_WORLD_OPEN", CRITICAL, name="sl-prod-open")
F_FORMULA = finding("OCI_SECLIST_SSH_WORLD_OPEN", CRITICAL, name='=HYPERLINK("http://x","y")')
F_BUCKET = finding("OCI_BUCKET_PUBLIC_ACCESS", HIGH, name="public-assets",
                   status=ACCEPTED_RISK, reason="Public website assets, approved by CISO",
                   by="alice")
F_ADMIN = finding("OCI_IAM_POLICY_TENANCY_ADMIN", CRITICAL, name="Tenant Admin Policy",
                  status=SUPPRESSED, reason="   ", by="bob")
F_DEFAULT_SL = finding("OCI_DEFAULT_SECLIST_WORLD_OPEN", LOW, name="Default Security List")
F_IMDS = finding("OCI_COMPUTE_LEGACY_IMDS_ENABLED", MEDIUM, name="vm-legacy", verdict=UNKNOWN)
FINDINGS = [F_SSH, F_FORMULA, F_BUCKET, F_ADMIN, F_DEFAULT_SL, F_IMDS]

RULE_STATS = {}
stats(RULE_STATS, "OCI_SECLIST_SSH_WORLD_OPEN", SAFE=8, VULNERABLE=2)
stats(RULE_STATS, "OCI_BUCKET_PUBLIC_ACCESS", SAFE=5)       # accepted one sits in SAFE
stats(RULE_STATS, "OCI_IAM_POLICY_TENANCY_ADMIN", SAFE=3)   # suppressed one sits in SAFE
stats(RULE_STATS, "OCI_DEFAULT_SECLIST_WORLD_OPEN", SAFE=9, VULNERABLE=1)
stats(RULE_STATS, "OCI_COMPUTE_LEGACY_IMDS_ENABLED", SAFE=4, UNKNOWN=1)

ACCOUNT = {"id": str(uuid.uuid4()), "name": "Prod Tenancy/1", "provider": PROVIDER,
           "environment": "Production", "provider_account_id": "ocid1.tenancy.oc1..test",
           "tenant_or_region": "ap-mumbai-1", "client": "KOEL",
           "last_discovery": (NOW - timedelta(hours=3)).isoformat()}
ASSESSMENT = SimpleNamespace(
    finished_at=NOW - timedelta(hours=1), outcome="PARTIAL", triggered_by="discovery:abc",
    coverage_pct=52.9, coverage={"network": 80.0, "identity": 40.0}, resources_evaluated=300,
    rules_evaluated=40, incomplete_reasons=["Audit log retention could not be read"])


def build_pack(framework=None, findings=FINDINGS, prepared_by="tester", include_raw=False):
    view = comp.build(findings=findings, rule_stats=RULE_STATS, account_id=ACCOUNT["id"],
                      framework=framework, provider=PROVIDER)
    return view, cx.assemble_pack(
        account=dict(ACCOUNT), view=view, findings=findings, rule_stats=RULE_STATS,
        assessment=ASSESSMENT, framework=framework, generated_at=NOW,
        prepared_by=prepared_by, include_raw=include_raw,
        assessment_stale_after_s=6 * 3600, inventory_stale_after_s=24 * 3600)


def control(pack, fw, cid):
    for f in pack["frameworks"]:
        if f["key"] == fw:
            return next((c for c in f["controls"] if c["control_id"] == cid), None)
    return None


def check_row(c, rule_id):
    return next((x for x in c["checks"] if x["rule_id"] == rule_id), None)


VIEW, PACK = build_pack()

# ═════════════════════════════════════════════════════════════════════════════
print("\n=== A. The pack is the Compliance tab, with its evidence ===")
# ═════════════════════════════════════════════════════════════════════════════
view_rows = {(r["framework"], r["control_id"]): r for r in VIEW["controls"]}
mismatch_state, mismatch_breach, seen = [], [], 0
for fw in PACK["frameworks"]:
    for c in fw["controls"]:
        seen += 1
        row = view_rows[(fw["key"], c["control_id"])]
        if c["state"] != row["state"]:
            mismatch_state.append((fw["key"], c["control_id"]))
        if len(c["evidence"]["breaching"]) != row["failing_findings"]:
            mismatch_breach.append((fw["key"], c["control_id"]))
eq("every control in the view is in the pack", seen, len(VIEW["controls"]))
eq("every control state equals the Compliance tab's", mismatch_state, [])
eq("every control's violation evidence is exactly its failing findings", mismatch_breach, [])
for fw in PACK["frameworks"]:
    s, v = fw["summary"], VIEW["frameworks"][fw["key"]]
    if (s["passed"], s["failed"], s["pass_rate"], s["coverage"]) != (
            v["controls_passed"], v["controls_failed"], v["control_pass_rate"],
            v["assessment_coverage"]):
        fail.append(f"summary {fw['key']}")
        print(f"  [FAIL] {fw['key']} summary differs from the view")
check("framework pass rate, coverage and counts equal the view's",
      not any(x.startswith("summary") for x in fail))
eq("frameworks follow the catalog order and the account's cloud",
   PACK["scope"]["frameworks"], [f for f in ctl.FRAMEWORKS if f in VIEW["frameworks"]])
check("CIS AWS / CIS Azure are not in an OCI pack",
      not {"CIS-AWS", "CIS-AZURE"} & set(PACK["scope"]["frameworks"]))

c21 = control(PACK, "CIS-OCI", "2.1")
eq("CIS OCI 2.1 failed by the world-open SSH lists", c21["state"], ctl.FAIL)
eq("2.1 evidence names both breaching findings",
   sorted(c21["evidence"]["breaching"]), sorted([F_SSH.fingerprint, F_FORMULA.fingerprint]))
ssh = check_row(c21, "OCI_SECLIST_SSH_WORLD_OPEN")
eq("2.1 check: evaluated 10, conforming 8, breaching 2",
   (ssh["evaluated"], ssh["conforming"], ssh["breaching"]), (10, 8, 2))

# ═════════════════════════════════════════════════════════════════════════════
print("\n=== B. Exceptions are never hidden ===")
# ═════════════════════════════════════════════════════════════════════════════
c511 = control(PACK, "CIS-OCI", "5.1.1")
eq("CIS OCI 5.1.1 reads PASS (the violation was accepted)", c511["state"], ctl.PASS)
check("…and is flagged as passed only with exceptions", c511["passed_with_exceptions"])
eq("5.1.1 lists the accepted finding as its exception",
   c511["evidence"]["exceptions"], [F_BUCKET.fingerprint])
bucket = check_row(c511, "OCI_BUCKET_PUBLIC_ACCESS")
eq("the accepted subject is subtracted from conforming (5 SAFE - 1 accepted)",
   (bucket["evaluated"], bucket["conforming"], bucket["exceptions"]), (5, 4, 1))
cc61 = control(PACK, "SOC2", "CC6.1")
check("SOC 2 CC6.1 is flagged too — the same accepted HIGH would fail it",
      cc61["passed_with_exceptions"], cc61["state"])
c12 = control(PACK, "CIS-OCI", "1.2")
check("a SUPPRESSED tenancy-admin policy is an exception on CIS OCI 1.2",
      c12["passed_with_exceptions"] and c12["evidence"]["exceptions"] == [F_ADMIN.fingerprint])

reg_by_fp = {e["fingerprint"]: e for e in PACK["exceptions"]}
eq("the register holds both decisions", sorted(reg_by_fp), sorted([F_BUCKET.fingerprint,
                                                                   F_ADMIN.fingerprint]))
eq("the accepted risk keeps its reason and decider",
   (reg_by_fp[F_BUCKET.fingerprint]["reason"], reg_by_fp[F_BUCKET.fingerprint]["decided_by"],
    reg_by_fp[F_BUCKET.fingerprint]["reason_missing"]),
   ("Public website assets, approved by CISO", "alice", False))
check("a whitespace-only reason counts as no reason", reg_by_fp[F_ADMIN.fingerprint]["reason_missing"])
check("the register cites the controls each decision affects",
      "CIS-OCI 5.1.1" in reg_by_fp[F_BUCKET.fingerprint]["controls"]
      and "SOC2 CC6.1" in reg_by_fp[F_BUCKET.fingerprint]["controls"])
eq("totals count exceptions and the reasonless one",
   (PACK["totals"]["exceptions"], PACK["totals"]["exceptions_without_reason"]), (2, 1))
eq("totals count every passed-with-exceptions control",
   PACK["totals"]["passed_with_exceptions"],
   sum(1 for fw in PACK["frameworks"] for c in fw["controls"] if c["passed_with_exceptions"]))
check("a muted finding is not a remediation priority",
      not any(p["rule_id"] in ("OCI_BUCKET_PUBLIC_ACCESS", "OCI_IAM_POLICY_TENANCY_ADMIN")
              for p in PACK["remediation_priorities"]))

# Without the exception, the same bucket fails the control: the flag is real.
_, open_pack = build_pack(findings=[F_SSH, F_FORMULA, F_DEFAULT_SL, F_IMDS,
                                    finding("OCI_BUCKET_PUBLIC_ACCESS", HIGH, name="b")])
eq("an OPEN public bucket fails 5.1.1 (control for the exception test)",
   control(open_pack, "CIS-OCI", "5.1.1")["state"], ctl.FAIL)

# ═════════════════════════════════════════════════════════════════════════════
print("\n=== C. Severity bars, undetermined evidence, gaps ===")
# ═════════════════════════════════════════════════════════════════════════════
c25 = control(PACK, "CIS-OCI", "2.5")
eq("a LOW default-security-list finding fails strict CIS OCI 2.5", c25["state"], ctl.FAIL)
iso89 = control(PACK, "ISO27001", "8.9")
default_sl = check_row(iso89, "OCI_DEFAULT_SECLIST_WORLD_OPEN")
check("…but is advisory, not a breach, against ISO 27001 8.9",
      F_DEFAULT_SL.fingerprint not in iso89["evidence"]["breaching"]
      and default_sl["advisory"] == 1 and default_sl["breaching"] == 0,
      f"state={iso89['state']} check={default_sl}")
c31 = control(PACK, "CIS-OCI", "3.1")
eq("an UNKNOWN IMDS result with passing subjects is PARTIAL", c31["state"], ctl.PARTIAL)
eq("…with the undetermined finding as its evidence",
   c31["evidence"]["undetermined"], [F_IMDS.fingerprint])
expected_ne = sorted((r["framework"], r["control_id"]) for r in VIEW["controls"]
                     if r["state"] == ctl.NOT_EVALUATED)
eq("every not-evaluated control is in the manual-evidence register",
   sorted((c["framework"], c["control_id"]) for c in PACK["not_evaluated"]), expected_ne)
check("no not-evaluated control is counted as passed",
      all(control(PACK, fw, cid)["state"] != ctl.PASS for fw, cid in expected_ne))
check("not-applicable controls carry their reason",
      all(c["reason"] for c in PACK["not_applicable"]))

cc61_rules = [x["rule_id"] for x in cc61["checks"]]
check("SOC 2 CC6.1 on OCI lists the OCI bucket check",
      "OCI_BUCKET_PUBLIC_ACCESS" in cc61_rules, cc61_rules)
check("…and none of the S3 / RDS / Azure checks that cannot run here",
      not any(r.startswith(("AWS_", "AZURE_")) for r in cc61_rules), cc61_rules)
top = PACK["remediation_priorities"][0]
eq("remediation starts with a CRITICAL", top["severity"], CRITICAL)
ssh_pri = next(p for p in PACK["remediation_priorities"]
               if p["rule_id"] == "OCI_SECLIST_SSH_WORLD_OPEN")
check("the SSH priority names CIS OCI 2.1 and counts both lists",
      "CIS-OCI 2.1" in ssh_pri["controls"] and ssh_pri["resources"] == 2, ssh_pri)
eq("the finding registry records each citation's role",
   {r["role"] for r in PACK["findings"][F_BUCKET.fingerprint]["controls"]}, {"exception"})
check("freshness is reported, and a 1-hour-old assessment is not stale",
      PACK["assessment"]["assessment_stale"] is False
      and PACK["assessment"]["inventory_stale"] is False)
eq("incomplete collection reasons are carried into the pack",
   PACK["assessment"]["incomplete_reasons"], ["Audit log retention could not be read"])

_, cis_pack = build_pack(framework="CIS-OCI")
eq("one framework narrows the pack", cis_pack["scope"]["frameworks"], ["CIS-OCI"])
eq("…and labels it by name", cis_pack["scope"]["label"], "CIS OCI")

# ═════════════════════════════════════════════════════════════════════════════
print("\n=== D. Integrity ===")
# ═════════════════════════════════════════════════════════════════════════════
check("the pack verifies", cx.verify_pack(PACK))
roundtrip = json.loads(cx.render_json(PACK))
check("…and still verifies after a JSON round trip (an auditor can recompute it)",
      cx.verify_pack(roundtrip))
eq("the hash is deterministic for the same inputs",
   build_pack()[1]["integrity"]["sha256"], PACK["integrity"]["sha256"])
tampered = json.loads(cx.render_json(PACK))
tampered["frameworks"][0]["summary"]["failed"] += 1
check("editing any value breaks verification", not cx.verify_pack(tampered))
check("a different preparer is a different pack",
      build_pack(prepared_by="someone else")[1]["integrity"]["sha256"]
      != PACK["integrity"]["sha256"])
eq("the pack ID is the hash's first 12 hex digits, upper-case",
   PACK["integrity"]["pack_id"], PACK["integrity"]["sha256"][:12].upper())

# ═════════════════════════════════════════════════════════════════════════════
print("\n=== E. Formats ===")
# ═════════════════════════════════════════════════════════════════════════════


def pdf_text(data: bytes) -> str:
    """Decode ReportLab's content streams (ASCII85 + Flate) to searchable text.
    The regex anchors on '>>' because 'endstream' itself contains 'stream'."""
    out = []
    for head, body in re.findall(rb"<<(.*?)>>\s*stream\r?\n(.*?)endstream", data, re.S):
        try:
            if b"ASCII85Decode" in head:
                body = base64.a85decode(body.strip().rstrip(b"~>").strip())
            if b"FlateDecode" in head:
                body = zlib.decompress(body)
        except Exception:  # noqa: BLE001
            continue
        out.append(body.decode("latin-1"))
    return re.sub(r"\\([()\\])", r"\1", "\n".join(out))


name, media, pdf = cx.render(PACK, "pdf")
eq("PDF media type", media, "application/pdf")
check("PDF bytes", pdf[:5] == b"%PDF-")
text = pdf_text(pdf)
check("every page names the pack", f"Pack {PACK['integrity']['pack_id']}" in text)
check("the PDF flags a control passed only with exceptions",
      "Passed only with exceptions" in text)
check("the PDF flags the reasonless exception", "No reason recorded" in text)
check("the PDF prints the full SHA-256", PACK["integrity"]["sha256"] in text)
check("the PDF carries a sign-off block", "Reviewed by" in text and "Approved by" in text)
check("the PDF names the failing control", "Ensure no security lists allow ingress" in text)

name, media, xlsx = cx.render(PACK, "xlsx")
wb = load_workbook(io.BytesIO(xlsx))
eq("workbook sheets", wb.sheetnames,
   ["Summary", "Controls", "Checks", "Evidence", "Exceptions", "Remediation",
    "Not Evaluated", "Not Applicable", "Assessment", "Method"])
eq("one Controls row per control", wb["Controls"].max_row - 1, PACK["totals"]["controls"])
refs = sum(len(c["evidence"][k]) for fw in PACK["frameworks"] for c in fw["controls"]
           for k in ("breaching", "undetermined", "exceptions"))
eq("one Evidence row per (control, finding)", wb["Evidence"].max_row - 1, refs)
eq("one Exceptions row per decision", wb["Exceptions"].max_row - 1, 2)
formula_cells = [c for row in wb["Evidence"].iter_rows(min_row=2) for c in row
                 if isinstance(c.value, str) and c.value.startswith("=HYPERLINK")]
check("a formula-shaped resource name is present…", bool(formula_cells))
check("…and stored as text, never as a formula",
      all(c.data_type == "s" for c in formula_cells), [c.data_type for c in formula_cells])
headers = [c.value for c in wb["Evidence"][1]]
check("no raw provider evidence unless asked",
      "Raw provider evidence (JSON)" not in headers)
_, raw_pack = build_pack(include_raw=True)
raw_wb = load_workbook(io.BytesIO(cx.render(raw_pack, "xlsx")[2]))
check("raw evidence column present when asked",
      "Raw provider evidence (JSON)" in [c.value for c in raw_wb["Evidence"][1]])
check("raw evidence is in the JSON finding rows when asked",
      all("evidence" in f for f in raw_pack["findings"].values())
      and not any("evidence" in f for f in PACK["findings"].values()))
summary_values = [c.value for row in wb["Summary"].iter_rows() for c in row]
check("the Summary sheet carries the pack hash", PACK["integrity"]["sha256"] in summary_values)

name, media, blob = cx.render(PACK, "zip")
z = zipfile.ZipFile(io.BytesIO(blob))
names = z.namelist()
eq("the ZIP holds PDF, XLSX, JSON, checksums and a README",
   sorted(n.rsplit(".", 1)[-1] for n in names), ["json", "pdf", "txt", "txt", "xlsx"])
sums = dict(line.split("  ", 1)[::-1] for line in
            z.read("SHA256SUMS.txt").decode().strip().splitlines())
check("every checksum matches its member",
      all(hashlib.sha256(z.read(n)).hexdigest() == h for n, h in sums.items())
      and len(sums) == 3, sums)
check("the README names the pack hash", PACK["integrity"]["sha256"] in z.read("README.txt").decode())
check("the JSON inside the ZIP verifies",
      cx.verify_pack(json.loads(z.read(next(n for n in names if n.endswith(".json"))))))

eq("filename: account slug, framework, UTC stamp",
   cx.filename(cis_pack, "pdf"), "compliance-audit-Prod-Tenancy-1-CIS-OCI-20261001-0930Z.pdf")
eq("filename without a framework says all-frameworks",
   cx.filename(PACK, "xlsx"), "compliance-audit-Prod-Tenancy-1-all-frameworks-20261001-0930Z.xlsx")
try:
    cx.render(PACK, "docx")
    check("an unknown format is refused", False)
except ValueError:
    check("an unknown format is refused", True)

# ═════════════════════════════════════════════════════════════════════════════
print("\n=== F. The route ===")
# ═════════════════════════════════════════════════════════════════════════════
from starlette.routing import Match  # noqa: E402

from app.api.routes import security_center as sc  # noqa: E402

acct = uuid.uuid4()
scope = {"type": "http", "method": "GET", "path_params": {},
         "path": f"/cloud/security-center/{acct}/compliance/export"}
matched = next((r for r in sc.router.routes if r.matches(scope)[0] == Match.FULL), None)
eq("/compliance/export reaches the export handler, not a control named 'export'",
   getattr(matched, "name", None), "export_compliance")

denials = []


async def _recorder(principal, permission, request):
    denials.append(permission)


rbac._record_denial = _recorder      # never write a real Event Log row from a test
REQ = SimpleNamespace(url=SimpleNamespace(path="/x"), method="GET")


def principal(perms):
    return rbac.Principal(name="t", role="test", permissions=set(perms), authenticated=True,
                          enforced=True)


def call(**kw):
    args = dict(account_id=acct, request=REQ, framework=None, fmt="pdf", raw=False, svc=None,
                principal=principal(rbac.ROLE_PERMISSIONS[rbac.ROLE_ADMIN]))
    args.update(kw)
    try:
        asyncio.run(sc.export_compliance(**args))
        return None
    except HTTPException as exc:
        return exc.status_code


eq("an unknown format is a 400", call(fmt="docx"), 400)
eq("an unknown framework is a 400", call(framework="SOX"), 400)
check("a viewer cannot export (no reports:export)",
      rbac.REPORT_EXPORT not in rbac.ROLE_PERMISSIONS[rbac.ROLE_VIEWER])
eq("…and is refused with a 403", call(principal=principal(rbac.ROLE_PERMISSIONS[rbac.ROLE_VIEWER])), 403)
eq("…which is audited as a reports:export denial", denials[-1:], [rbac.REPORT_EXPORT])
eq("raw evidence without security:view_raw is a 403",
   call(raw=True, principal=principal({rbac.SECURITY_VIEW, rbac.REPORT_EXPORT})), 403)
eq("…audited as a security:view_raw denial", denials[-1:], [rbac.SECURITY_VIEW_RAW])

print(f"\n{len(ok)} passed, {len(fail)} failed")
if fail:
    print("FAILED:", *fail, sep="\n  - ")
sys.exit(1 if fail else 0)
