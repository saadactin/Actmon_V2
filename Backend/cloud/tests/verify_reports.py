"""The resource assessment report builder: scope cap, section toggles, every
renderer's structural honesty. No database, no ASGI -- where a production
function reads the database, it is given a scripted fake session, so the
real function runs rather than a re-implementation of its rule.

What this suite pins, and why each matters:

* **The resources-mode cap is enforced with zero DB queries.** An explicit
  resource_ids list over REPORT_MAX_RESOURCES must raise before touching the
  database -- and say the report's cap (1,000), not the old export's (200).
* **A disabled section produces no rows, in every format.**
* **NA is never zero.** A resource with no stored cost renders as blank /
  "Not available" everywhere, never 0.
* **A single very long finding detail does not break the PDF.**
* **Excel cannot store a timezone-aware datetime.**
* **A CSV cell starting with =+-@ is neutralised.**
* **The reporting period is real.** Presets resolve to whole days in the
  requested time zone; sub-day presets that nothing could honour are refused;
  spend anomalies outside the period are not reported.
* **Every section prints its own columns.** The PDF Cost section printed no
  amounts and Inventory no type/region while every section went through one
  findings-shaped table.
* **No filler rows; account-level findings are reported.** Security matching
  is case-insensitive (Azure), and findings tied to no resource appear in a
  whole-account report and are counted, not silently dropped, in a scoped one.
* **Detail level, grouping and sorting change the output in every format.**
"""
import pathlib
import re
import sys
import uuid
import zlib
from datetime import date, datetime, timezone
from types import SimpleNamespace

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))

import asyncio  # noqa: E402
import csv  # noqa: E402
import io  # noqa: E402

from pydantic import ValidationError  # noqa: E402

from app.schemas.reports import (  # noqa: E402
    ReportDateRange, ReportLayout, ReportScope,
)
from app.services import report_service as rs  # noqa: E402
from app.services import resource_assessment_service as ras  # noqa: E402
from app.services.report_renderers import (  # noqa: E402
    csv_renderer, json_renderer, pdf_renderer, xlsx_renderer,
)
from app.services.resource_assessment_service import TooManyResourcesError  # noqa: E402
from openpyxl import load_workbook  # noqa: E402


def _parse_csv_rows(text: str):
    """[{header: value}, ...], tolerating the leading BOM and the blank-line
    summary block the renderer appends after the tabular rows."""
    reader = csv.reader(io.StringIO(text.lstrip("﻿")))
    header = next(reader)
    rows = []
    for raw in reader:
        if not raw or len(raw) != len(header):
            break  # the summary block starts here
        rows.append(dict(zip(header, raw)))
    return rows


def _pdf_text(data: bytes) -> str:
    """Every string a PDF draws. ReportLab writes content streams as
    ASCII85-encoded Flate, so both layers are undone before reading the Tj
    operators."""
    import base64

    out = []
    # Anchored on the dictionary's ">>": a bare "stream\n" also matches the
    # tail of "endstream\n" and would swallow the next object whole.
    for raw in re.findall(rb">>\s*stream\r?\n(.*?)endstream", data, re.S):
        raw = raw.strip()
        try:
            if raw.endswith(b"~>"):
                raw = base64.a85decode(raw[:-2].replace(b"\n", b"").replace(b"\r", b""))
            out.append(zlib.decompress(raw).decode("latin-1"))
        except (ValueError, zlib.error):
            continue  # fonts, images: not page text
    # PDF string literals escape their own parentheses: "(Region: r1 \(1\)) Tj".
    strings = re.findall(r"\(((?:\\.|[^\\)])*)\)\s*Tj", "\n".join(out))
    return "\n".join(re.sub(r"\\(.)", r"\1", s) for s in strings)


ok, fail = [], []


def check(label, cond, detail=""):
    (ok if cond else fail).append(label)
    print(f"  [{'PASS' if cond else 'FAIL'}] {label}" + (f"  {detail}" if detail else ""))


def eq(label, got, want):
    check(label, got == want, f"got {got!r}, want {want!r}")


def _dataset(sections: dict, *, notes=None, coverage=None, **extra) -> dict:
    return {
        "schema_version": 2,
        "account": {"id": "a", "name": "zz-test-acct", "provider": "AWS"},
        "scope": {"mode": "account"},
        "period": {"preset": "30d", "label": "Last 30 days", "start": "2026-09-22",
                   "end": "2026-10-21", "days": 30, "timezone": "UTC"},
        "resources_requested": 1,
        "resources_resolved": 1,
        "resources_missing": [],
        "notes": notes or [],
        "coverage": coverage or {},
        "freshness": {"warnings": []},
        "sections": sections,
        **extra,
    }


SECURITY_ROW = {
    "account_name": "zz-test-acct", "provider": "AWS", "region": "ap-south-1",
    "resource_name": "r1", "resource_type": "S3Bucket", "provider_resource_id": "p1",
    "assessment_category": "Security", "severity": "CRITICAL", "title": "Bad thing",
    "detail": "detail text", "metric_value": None, "status": "OPEN",
    "detected_at": "2026-09-10T10:30:49.231166+00:00", "attribution": "Resource",
}
COST_ROW = {
    "resource_name": "bigdb", "provider_resource_id": "ocid1.db", "resource_type": "DbSystem",
    "region": "ap-mumbai-1", "cost_amount": 1309674.22, "currency": "INR",
    "cost_period_days": 30, "billed_service": "Database", "status": "Billed (provider billing)",
}
COST_ROW_NA = {
    "resource_name": "r1", "provider_resource_id": "p1", "cost_amount": None,
    "currency": None, "cost_period_days": 30, "status": "Billed, amount not reported",
}
INVENTORY_ROW = {
    "resource_name": "r1", "provider_resource_id": "p1", "resource_type": "S3Bucket",
    "region": "ap-south-1", "status": "UNKNOWN", "environment": None, "owner": None,
    "resource_group": None, "compartment": None, "tags": {},
    "discovered_at": "2026-08-06T06:11:55.903227+00:00",
    "last_seen_at": "2026-10-12T11:56:47.990103+00:00",
}

# ═════════════════════════════════════════════════════════════════════════════
print("\n=== A. Schema validation ===")
# ═════════════════════════════════════════════════════════════════════════════
try:
    ReportScope(mode="resources", resource_ids=[])
    check("empty resource_ids rejected for mode=resources", False)
except ValidationError:
    check("empty resource_ids rejected for mode=resources", True)

check("mode='account' needs no resource_ids",
      bool(ReportScope(mode="account")))

for bad, label in ((dict(format="docx"), "unknown format rejected"),
                   (dict(group_by="owner"), "unknown group_by rejected"),
                   (dict(sort_by="random"), "unknown sort_by rejected")):
    try:
        ReportLayout(**bad)
        check(label, False)
    except ValidationError:
        check(label, True)

eq("legacy detail level 'technical' is read as 'detailed', not refused",
   ReportLayout(detail_level="technical").detail_level, "detailed")
eq("an empty group_by from a <select> means no grouping",
   ReportLayout(group_by="").group_by, None)

for bad, label in ((dict(preset="custom"), "custom preset without start/end rejected"),
                   (dict(preset="1h"), "sub-day preset '1h' rejected (nothing is hourly)"),
                   (dict(preset="custom", start="2026-10-10", end="2026-10-01"),
                    "custom start after end rejected"),
                   (dict(preset="custom", start="2025-01-01", end="2026-10-01"),
                    "custom range over 366 days rejected"),
                   (dict(timezone="Mars/Olympus"), "unknown time zone rejected")):
    try:
        ReportDateRange(**bad)
        check(label, False)
    except ValidationError:
        check(label, True)

legacy = ReportDateRange(preset="custom", start="2026-10-01T09:30", end="2026-10-05T18:00")
eq("an old datetime-local draft value keeps its date", (legacy.start, legacy.end),
   (date(2026, 10, 1), date(2026, 10, 5)))
eq("a custom range counts both ends", legacy.resolved_days(), 5)

# 2026-10-20 20:00 UTC is already 21 Oct in Kolkata (+05:30).
NOW = datetime(2026, 10, 20, 20, 0, tzinfo=timezone.utc)
eq("'7d' resolves to 7 whole days ending today in the requested zone",
   ReportDateRange(preset="7d", timezone="Asia/Kolkata").resolve(NOW),
   (date(2026, 10, 15), date(2026, 10, 21)))
eq("...and in UTC, today is still the 20th",
   ReportDateRange(preset="7d", timezone="UTC").resolve(NOW),
   (date(2026, 10, 14), date(2026, 10, 20)))
eq("'mtd' starts on the 1st", ReportDateRange(preset="mtd").resolve(NOW),
   (date(2026, 10, 1), date(2026, 10, 20)))
eq("'last_month' is the whole previous calendar month",
   ReportDateRange(preset="last_month").resolve(NOW), (date(2026, 9, 1), date(2026, 9, 30)))
eq("...and is labelled by its month", ReportDateRange(preset="last_month").describe(NOW)["label"],
   "September 2026")

# ═════════════════════════════════════════════════════════════════════════════
print("\n=== B. Scope cap enforced before any DB query ===")
# ═════════════════════════════════════════════════════════════════════════════
too_many = ReportScope(mode="resources", resource_ids=[uuid.uuid4() for _ in range(rs.REPORT_MAX_RESOURCES + 1)])
try:
    asyncio.run(rs.resolve_report_scope(None, uuid.uuid4(), too_many))  # db=None: must not be touched
    check(f"{rs.REPORT_MAX_RESOURCES + 1} explicit resource_ids raises TooManyResourcesError", False)
except TooManyResourcesError as exc:
    check(f"{rs.REPORT_MAX_RESOURCES + 1} explicit resource_ids raises TooManyResourcesError", True)
    check("...and the message names the report's own cap, not the export's 200",
          f"capped at {rs.REPORT_MAX_RESOURCES}" in str(exc), str(exc))

dupe_ids = [uuid.uuid4()] * 3
ids, total = asyncio.run(rs.resolve_report_scope(None, uuid.uuid4(), ReportScope(mode="resources", resource_ids=dupe_ids)))
eq("duplicate resource_ids are de-duplicated", len(ids), 1)
eq("...and the resolved total reflects that", total, 1)

# ═════════════════════════════════════════════════════════════════════════════
print("\n=== C. A disabled section is absent from every format ===")
# ═════════════════════════════════════════════════════════════════════════════
dataset_security_only = _dataset({"security": [SECURITY_ROW]})
csv_text = csv_renderer.build_csv(dataset_security_only)
csv_rows = _parse_csv_rows(csv_text)
eq("CSV has exactly the one enabled section's row", len(csv_rows), 1)
eq("...tagged with the Security section label", csv_rows[0]["Section"], "Security")
check("...and no row is tagged Inventory or Cost",
      all(r["Section"] not in ("Inventory", "Cost") for r in csv_rows))

json_text = json_renderer.build_json(dataset_security_only)
check("JSON omits a disabled section entirely", '"cost"' not in json_text and '"inventory"' not in json_text)

wb_bytes = xlsx_renderer.build_xlsx(dataset_security_only)
wb = load_workbook(io.BytesIO(wb_bytes))
eq("XLSX has exactly Summary + the one enabled section's sheet",
   set(wb.sheetnames), {"Summary", "Security Findings"})

pdf_bytes = pdf_renderer.build_pdf(dataset_security_only, detail_level="standard")
check("PDF is a valid, non-trivial file", pdf_bytes[:4] == b"%PDF" and len(pdf_bytes) > 500)

# ═════════════════════════════════════════════════════════════════════════════
print("\n=== D. NA is never zero ===")
# ═════════════════════════════════════════════════════════════════════════════
dataset_cost = _dataset({"cost": [COST_ROW_NA], "inventory": [INVENTORY_ROW]})
cost_rows = _parse_csv_rows(csv_renderer.build_csv(dataset_cost))
cost_row = next(r for r in cost_rows if r["Section"] == "Cost")
eq("CSV's cost cell is blank, not '0', for a resource with no stored amount",
   cost_row["Billed Cost"], "")

wb2 = load_workbook(io.BytesIO(xlsx_renderer.build_xlsx(dataset_cost)))
cost_ws = wb2["Cost"]
header = [c.value for c in cost_ws[1]]
cost_col = header.index("Billed Cost") + 1
check("XLSX cost cell is None, not 0, for a resource with no stored amount",
      cost_ws.cell(row=2, column=cost_col).value is None)
check("PDF prints 'Not available', never 0.00, for it",
      "Not available" in _pdf_text(pdf_renderer.build_pdf(dataset_cost)))

# ═════════════════════════════════════════════════════════════════════════════
print("\n=== E. A single very long finding detail does not break the PDF ===")
# ═════════════════════════════════════════════════════════════════════════════
long_row = {**SECURITY_ROW, "detail": "x" * 5845}
dataset_long = _dataset({"security": [long_row] * 50})  # 50 rows, one page-breaking cell each
try:
    pdf_bytes2 = pdf_renderer.build_pdf(dataset_long, detail_level="standard")
    check("PDF renders 50 rows with a 5,845-char detail field each", pdf_bytes2[:4] == b"%PDF")
except Exception as exc:  # noqa: BLE001
    check("PDF renders 50 rows with a 5,845-char detail field each", False, str(exc))

# ═════════════════════════════════════════════════════════════════════════════
print("\n=== F. Excel cannot store a timezone-aware datetime -- stripped, not fatal ===")
# ═════════════════════════════════════════════════════════════════════════════
try:
    xlsx_renderer.build_xlsx(_dataset({"inventory": [INVENTORY_ROW]}))
    check("tz-aware discovered_at/last_seen_at do not raise", True)
except TypeError as exc:
    check("tz-aware discovered_at/last_seen_at do not raise", False, str(exc))

# ═════════════════════════════════════════════════════════════════════════════
print("\n=== G. CSV formula-injection guard ===")
# ═════════════════════════════════════════════════════════════════════════════
injected_row = {**SECURITY_ROW, "title": "=1+1", "detail": "@SUM(A1:A9)"}
csv_injected = csv_renderer.build_csv(_dataset({"security": [injected_row]}))
injected_rows = _parse_csv_rows(csv_injected)
eq("a '=' title is prefixed with a guard quote", injected_rows[0]["Title"], "'=1+1")
eq("a '@' detail is prefixed with a guard quote", injected_rows[0]["Detail"], "'@SUM(A1:A9)")
check("CSV carries a UTF-8 BOM", csv_injected.startswith("﻿"))
eq("a text value starting with '-' is still guarded", csv_renderer._cell("-2+3+cmd|' /C calc'!A0"),
   "'-2+3+cmd|' /C calc'!A0")
eq("...but a real negative number stays a number", csv_renderer._cell(-907.2), "-907.2")

# ═════════════════════════════════════════════════════════════════════════════
print("\n=== H. Every section prints its own columns ===")
# ═════════════════════════════════════════════════════════════════════════════
pdf_txt = _pdf_text(pdf_renderer.build_pdf(_dataset({"cost": [COST_ROW], "inventory": [INVENTORY_ROW]})))
check("PDF Cost section prints the billed amount and currency", "1,309,674.22 INR" in pdf_txt)
check("...and the billed service", "Database" in pdf_txt)
check("PDF Inventory section prints the type and region", "S3Bucket" in pdf_txt and "ap-south-1" in pdf_txt)
check("PDF header prints the reporting period", "2026-09-22" in pdf_txt and "2026-10-21" in pdf_txt)

wb3 = load_workbook(io.BytesIO(xlsx_renderer.build_xlsx(_dataset({"cost": [COST_ROW], "inventory": [INVENTORY_ROW]}))))
inv_header = [c.value for c in wb3["Inventory"][1]]
check("XLSX Inventory sheet has no empty Severity/Billed Cost columns",
      "Severity" not in inv_header and "Billed Cost" not in inv_header, str(inv_header))
cost_header = [c.value for c in wb3["Cost"][1]]
eq("XLSX Cost sheet stores the amount as a number",
   wb3["Cost"].cell(row=2, column=cost_header.index("Billed Cost") + 1).value, 1309674.22)
check("XLSX Cost sheet carries the currency", "Currency" in cost_header)

warn_ds = _dataset({"security": [SECURITY_ROW]},
                   freshness={"warnings": ["The last refresh finished PARTIAL: cost failed."],
                              "refresh_status": "PARTIAL", "last_refreshed_at": "2026-10-15T08:13:00+00:00"})
check("CSV report block carries the period and the freshness warning",
      "2026-09-22" in csv_renderer.build_csv(warn_ds)
      and "finished PARTIAL" in csv_renderer.build_csv(warn_ds))
check("PDF first page carries the freshness warning",
      "finished PARTIAL" in _pdf_text(pdf_renderer.build_pdf(warn_ds)))

# ═════════════════════════════════════════════════════════════════════════════
print("\n=== I. Sorting, grouping and detail level ===")
# ═════════════════════════════════════════════════════════════════════════════
rows = [
    {**SECURITY_ROW, "resource_name": "b", "severity": "LOW", "region": "r2"},
    {**SECURITY_ROW, "resource_name": "a", "severity": "CRITICAL", "region": "r2"},
    {**SECURITY_ROW, "resource_name": "c", "severity": "HIGH", "region": "r1"},
]
ds = _dataset({"security": [dict(r) for r in rows]})
rs.order_rows(ds, None, None)
eq("security defaults to most severe first", [r["severity"] for r in ds["sections"]["security"]],
   ["CRITICAL", "HIGH", "LOW"])
ds = _dataset({"security": [dict(r) for r in rows]})
rs.order_rows(ds, "region", "severity")
eq("grouped by region, severity order kept inside each group",
   [(r["region"], r["severity"]) for r in ds["sections"]["security"]],
   [("r1", "HIGH"), ("r2", "CRITICAL"), ("r2", "LOW")])

costs = [{**COST_ROW, "resource_name": "x", "cost_amount": None},
         {**COST_ROW, "resource_name": "y", "cost_amount": 5.0},
         {**COST_ROW, "resource_name": "z", "cost_amount": 50.0}]
ds = _dataset({"cost": costs})
rs.order_rows(ds, None, "cost")
eq("cost sorts highest first, and a missing amount goes last (never as zero)",
   [r["resource_name"] for r in ds["sections"]["cost"]], ["z", "y", "x"])
ds = _dataset({"inventory": [dict(INVENTORY_ROW, resource_name="b"), dict(INVENTORY_ROW, resource_name="a")]})
rs.order_rows(ds, None, "severity")
eq("a sort field a section lacks falls back to its natural order",
   [r["resource_name"] for r in ds["sections"]["inventory"]], ["a", "b"])

grouped_pdf = _pdf_text(pdf_renderer.build_pdf(
    {**_dataset({"security": [dict(r) for r in rows]}), "layout": {"group_by": "region"}}))
check("PDF prints a sub-heading per group", "Region: r1 (1)" in grouped_pdf and "Region: r2 (2)" in grouped_pdf)

without = [{"resource_name": "quiet", "provider_resource_id": "q", "resource_type": "VCN",
            "region": "r1", "reason": "No open security findings."}]


def _cov_ds():
    return _dataset({"security": [dict(SECURITY_ROW)], "executive_summary": {"resources_in_scope": 2}},
                    coverage={"security": {"status": "ok", "resources_checked": 2,
                                           "resources_with_data": 1, "resources_without_data": 1,
                                           "without": list(without)}})


ds = _cov_ds()
rs.apply_detail_level(ds, "executive")
check("executive drops every per-resource row section", "security" not in ds["sections"])
check("...keeps the summary", "executive_summary" in ds["sections"])
ds = _cov_ds()
rs.apply_detail_level(ds, "standard")
check("standard keeps finding rows", len(ds["sections"]["security"]) == 1)
check("...but only counts the resources with nothing to report",
      ds["coverage"]["security"]["without"] == [] and ds["coverage"]["security"]["resources_without_data"] == 1)
eq("...so its CSV has no coverage rows",
   [r["Section"] for r in _parse_csv_rows(csv_renderer.build_csv(ds))], ["Security"])
ds = _cov_ds()
rs.apply_detail_level(ds, "detailed")
eq("detailed lists them, with the reason, in the CSV",
   [(r["Section"], r["Resource"], r["Detail"]) for r in _parse_csv_rows(csv_renderer.build_csv(ds))][1],
   ("Coverage: Security", "quiet", "No open security findings."))
check("...in a Coverage sheet in Excel", "Coverage" in load_workbook(io.BytesIO(xlsx_renderer.build_xlsx(ds))).sheetnames)
check("...and after the section's table in the PDF",
      "Nothing to report (1)" in _pdf_text(pdf_renderer.build_pdf(ds)))

# ═════════════════════════════════════════════════════════════════════════════
print("\n=== J. Security: no filler, case-insensitive match, account-level findings ===")
# ═════════════════════════════════════════════════════════════════════════════


class _Result:
    def __init__(self, rows):
        self._rows = rows

    def scalars(self):
        return self

    def all(self):
        return list(self._rows)


class FakeDB:
    """Answers each `execute` with the next scripted result, in call order."""

    def __init__(self, *results):
        self._results = list(results)

    async def execute(self, _stmt):
        return _Result(self._results.pop(0))


ACCT = uuid.uuid4()


def _res(name, key):
    return SimpleNamespace(id=uuid.uuid4(), account_id=ACCT, resource_name=name,
                           provider_resource_id=key, resource_type="VirtualMachine",
                           region_or_zone="centralindia")


def _finding(key, title, severity="HIGH", resource_id=None):
    return SimpleNamespace(resource_key=key, resource_id=resource_id, title=title, severity=severity,
                           description="d", status="OPEN", first_detected_at=None,
                           resource_name=None, resource_type="RoleAssignment", region=None)


vm = _res("vm1", "/subscriptions/S/resourceGroups/RG/providers/Microsoft.Compute/virtualMachines/vm1")
quiet = _res("vm2", "/subscriptions/S/resourceGroups/RG/providers/Microsoft.Compute/virtualMachines/vm2")
outside = _res("vm3", "/subscriptions/S/resourceGroups/RG/providers/Microsoft.Compute/virtualMachines/vm3")
findings = [
    _finding(vm.provider_resource_id.lower(), "Lower-cased id from another Azure API"),
    _finding("/subscriptions/S/providers/Microsoft.Authorization/roleAssignments/x",
             "Owner over the entire subscription", "CRITICAL"),
    _finding(None, "Resources are missing ownership tags", "LOW"),
    _finding(outside.provider_resource_id, "On a resource outside the scope"),
]
inventory_rows = [(r.id, r.provider_resource_id) for r in (vm, quiet, outside)]


def _run_security(scope_mode, resources):
    out = {"sections": {}, "coverage": {}, "notes": []}
    bases = {r.id: ras._base_row(r, "acct", "Azure") for r in resources}
    asyncio.run(rs._security_section(FakeDB(findings, inventory_rows), out, ACCT, "acct", "Azure",
                                     resources, bases, scope_mode))
    return out


whole = _run_security("account", [vm, quiet, outside])
titles = [r["title"] for r in whole["sections"]["security"]]
check("a finding whose id differs only in case is matched to its resource",
      "Lower-cased id from another Azure API" in titles)
check("no 'No findings' filler row for the resource without findings",
      all(r["title"] != "No findings" for r in whole["sections"]["security"]))
eq("...it is counted in coverage instead", whole["coverage"]["security"]["resources_without_data"], 1)
check("whole-account: subscription-wide and keyless findings are included",
      "Owner over the entire subscription" in titles and "Resources are missing ownership tags" in titles)
eq("...labelled account-level", sorted(r["attribution"] for r in whole["sections"]["security"]
                                       if r["title"] in ("Owner over the entire subscription",
                                                         "Resources are missing ownership tags")),
   ["Account-level", "Account-level"])
eq("...and counted", whole["coverage"]["security"]["account_level_findings"], 2)

scoped = _run_security("filtered", [vm, quiet])
scoped_titles = [r["title"] for r in scoped["sections"]["security"]]
check("filtered: account-level findings are left out of the rows",
      "Owner over the entire subscription" not in scoped_titles)
check("filtered: a finding on an inventoried resource outside the scope is left out",
      "On a resource outside the scope" not in scoped_titles)
check("...and the report says how many account-level findings it left out",
      any("2 open account-level" in n for n in scoped["notes"]), str(scoped["notes"]))

# ═════════════════════════════════════════════════════════════════════════════
print("\n=== K. Spend anomalies follow the reporting period ===")
# ═════════════════════════════════════════════════════════════════════════════
from datetime import timedelta  # noqa: E402

days = [date(2026, 9, 1) + timedelta(days=i) for i in range(45)]
spikes = {date(2026, 9, 20), date(2026, 10, 10)}  # one in the lead-in, one in the period
series = [(d, "INR", 10000.0 if d in spikes else 100.0, 1) for d in days]
period_rows = asyncio.run(ras._account_anomaly_rows(
    FakeDB(series), uuid.uuid4(), (date(2026, 10, 1), date(2026, 10, 15))))
eq("only the spike inside the period is reported",
   [r["title"] for r in period_rows], ["Spend anomaly on 2026-10-10"])
clean = asyncio.run(ras._account_anomaly_rows(
    FakeDB([(d, "INR", 100.0, 1) for d in days]), uuid.uuid4(),
    (date(2026, 10, 1), date(2026, 10, 15))))
check("a clean period says so, naming the period, not '90 days'",
      clean[0]["title"] == "No anomalies" and "2026-10-01" in clean[0]["detail"], clean[0]["detail"])

# ═════════════════════════════════════════════════════════════════════════════
print("\n=== L. Spend change is like for like, or not at all ===")
# ═════════════════════════════════════════════════════════════════════════════
P = (date(2026, 10, 1), date(2026, 10, 10))  # a 10-day period
prev = [{"date": (date(2026, 9, 21) + timedelta(days=i)).isoformat(), "currency": "INR", "amount": 100.0}
        for i in range(10)]
cur = [{"date": (date(2026, 10, 1) + timedelta(days=i)).isoformat(), "currency": "INR", "amount": 110.0}
       for i in range(7)]  # billed through the 7th only
ch = rs.spend_change(prev + cur, P, "2026-10-07")
eq("billing lag cuts the current window at the last billed day",
   ch["current_window"], ["2026-10-01", "2026-10-07"])
eq("...and the previous window gets the same length, just before the period",
   ch["previous_window"], ["2026-09-24", "2026-09-30"])
eq("7 days at 110 against 7 days at 100 is +10%", ch["by_currency"]["INR"]["percent"], 10.0)
check("the old naive comparison (7 billed days vs 10) would have shown a fall instead",
      sum(p["amount"] for p in cur) < sum(p["amount"] for p in prev))
gap = [p for p in prev if p["date"] != "2026-09-27"]
eq("a missing day in either window makes it not comparable",
   rs.spend_change(gap + cur, P, "2026-10-07")["by_currency"]["INR"]["kind"], "not_comparable")
usd = [{**p, "currency": "USD"} for p in cur]
two = rs.spend_change(prev + usd, P, "2026-10-07")["by_currency"]
check("currencies are compared separately, never across",
      two["USD"]["kind"] == "not_comparable" and two["INR"]["kind"] == "not_comparable")
eq("no billing at all means no comparison", rs.spend_change([], P, None), None)

# ═════════════════════════════════════════════════════════════════════════════
print("\n=== M. Savings, top risks, posture, identity labels, period bounds ===")
# ═════════════════════════════════════════════════════════════════════════════
opt = lambda amount, cur="INR": SimpleNamespace(estimated_savings=amount, savings_currency=cur)  # noqa: E731
sv = rs.savings_summary([opt(100.0), opt(50.5), opt(None), opt(20.0, "USD"), opt(None, None)])
eq("savings are summed per currency", sv["monthly_by_currency"], {"INR": 150.5, "USD": 20.0})
eq("...and an unquantified saving is counted, never added as 0", (sv["priced_findings"], sv["unpriced_findings"]), (3, 2))

risk_rows = ([{**SECURITY_ROW, "title": "Open SSH", "severity": "CRITICAL", "remediation": "close it"}] * 3
             + [{**SECURITY_ROW, "title": "Open RDP", "severity": "CRITICAL"}]
             + [{**SECURITY_ROW, "title": "No tags", "severity": "LOW"}] * 9)
risks = rs._top_risks(risk_rows)
eq("one entry per rule, most severe first, then most resources",
   [(r["title"], r["resources"]) for r in risks], [("Open SSH", 3), ("Open RDP", 1), ("No tags", 9)])
eq("...carrying the rule's remediation", risks[0]["remediation"], "close it")


class _One:
    def __init__(self, row):
        self.row = row

    def scalars(self):
        return self

    def first(self):
        return self.row

    def scalar_one_or_none(self):
        return self.row


class OneDB:
    def __init__(self, row):
        self.row = row

    async def execute(self, _stmt):
        return _One(self.row)


withheld = SimpleNamespace(score=62, grade="C", score_reliable=False, coverage_pct=18.0,
                           finished_at=None, score_breakdown={})
ps = asyncio.run(rs._posture_score(OneDB(withheld), uuid.uuid4()))
eq("a withheld score is not reported as its raw number", ps["score"], None)
check("...it says it is withheld", ps["reliable"] is False and "withholds" in ps["note"])
eq("no assessment reads 'Not assessed', not 0",
   asyncio.run(rs._posture_score(OneDB(None), uuid.uuid4()))["grade"], "Not assessed")
withheld_pdf = _pdf_text(pdf_renderer.build_pdf(_dataset(
    {"executive_summary": {"resources_in_scope": 1, "posture": ps}})))
check("the PDF prints 'Withheld', not 62", "Withheld" in withheld_pdf and "62/100" not in withheld_pdf)

iid = uuid.uuid4()
labels = asyncio.run(rs._identity_labels(FakeDB([(iid, None, None, "AZURE_USER", "2e5a")]), [iid]))
eq("an identity Graph could not name is labelled as unresolved, not blank",
   labels[str(iid)], "AZURE_USER 2e5a (name not resolved)")

start, end = rs._period_bounds_utc((date(2026, 10, 1), date(2026, 10, 1)), "Asia/Kolkata")
eq("a Kolkata day starts at 18:30 UTC the day before", start.isoformat(), "2026-09-30T18:30:00+00:00")
eq("...and runs a full 24 hours", (end - start).total_seconds(), 86400.0)

# ═════════════════════════════════════════════════════════════════════════════
print("\n=== N. New sections in every format ===")
# ═════════════════════════════════════════════════════════════════════════════
CERT_ROW = {"resource_name": "koel.co.in", "resource_type": "CERTIFICATE", "title": "Expired",
            "status": "expired", "expires_at": "2024-04-28T00:00:00+00:00", "days_left": -907.2,
            "detail": "lb01 | *.koel.co.in", "in_use": True, "store": "OCI Certificates"}
CHANGE_ROW = {"resource_name": "nsg-web", "resource_type": "NSG", "region": "centralindia",
              "provider_resource_id": "/x/nsg-web", "title": "Modified: rules", "severity": "HIGH",
              "detail": "Ingress opened to 0.0.0.0/0", "metric_value": "a -> b", "status": "MORE_OPEN",
              "detected_at": "2026-10-13T00:00:00+00:00", "change_type": "MODIFIED"}
POSTURE = {"score": {"status": "ok", "score": 15, "grade": "F", "reliable": True, "coverage_pct": 52.9,
                     "assessed_at": "2026-10-22T10:55:49+00:00"},
           "frameworks_status": "ok",
           "frameworks": [{"framework": "CIS OCI", "controls_assessed": 12, "passed": 4, "failed": 8,
                           "partial": 0, "unknown": 0, "not_applicable": 0, "pass_rate": 33.3},
                          {"framework": "CIS Azure", "controls_assessed": 0, "passed": 0, "failed": 0,
                           "partial": 0, "unknown": 0, "not_applicable": 0, "pass_rate": None}]}
new_ds = _dataset({"security": [{**SECURITY_ROW, "remediation": "Run:\n  oci network ..."}],
                   "certificates": [CERT_ROW], "changes": [CHANGE_ROW], "posture": POSTURE})
rows_csv = _parse_csv_rows(csv_renderer.build_csv(new_ds))
eq("CSV carries the remediation", rows_csv[0]["Remediation"], "Run:\n  oci network ...")
cert_csv = next(r for r in rows_csv if r["Section"] == "Expiring Certificates & Secrets")
eq("CSV carries expiry and days left", (cert_csv["Expires At"][:10], cert_csv["Days Left"]), ("2024-04-28", "-907.2"))
check("CSV has a Configuration Changes row", any(r["Section"] == "Configuration Changes" for r in rows_csv))
wb4 = load_workbook(io.BytesIO(xlsx_renderer.build_xlsx(new_ds)))
check("Excel puts compliance frameworks on their own sheet", "Compliance" in wb4.sheetnames, str(wb4.sheetnames))
check("...with expiring credentials and config changes as sheets",
      {"Expiring Credentials", "Config Changes"} <= set(wb4.sheetnames))
eq("Excel stores days left as a number",
   wb4["Expiring Credentials"].cell(row=2, column=[c.value for c in wb4["Expiring Credentials"][1]].index("Days Left") + 1).value,
   -907.2)
new_pdf = _pdf_text(pdf_renderer.build_pdf({**new_ds, "layout": {"detail_level": "detailed"}}))
check("PDF prints the overdue certificate as days ago", "907 ago" in new_pdf)
check("PDF prints a framework that rated nothing as 'not rated', not 0%", "not rated" in new_pdf)
check("PDF prints the fix at the Detailed level", "How to fix:" in new_pdf)
check("...and not at Standard", "How to fix:" not in _pdf_text(pdf_renderer.build_pdf(new_ds)))

from app.api.routes.reports import report_slug  # noqa: E402

eq("filenames carry the account", report_slug("Azure - TOR (7e4839cb)"), "azure-tor-7e4839cb")
eq("...and never come out empty", report_slug("???"), "account")

# ═════════════════════════════════════════════════════════════════════════════
print("\n=== Summary ===")
# ═════════════════════════════════════════════════════════════════════════════
print(f"\n{len(ok)} passed, {len(fail)} failed")
if fail:
    print("FAILED:")
    for f in fail:
        print(f"  - {f}")
sys.exit(1 if fail else 0)
