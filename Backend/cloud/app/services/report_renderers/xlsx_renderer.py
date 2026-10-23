"""Excel rendering: one worksheet per enabled section plus a Summary sheet.

Only a worksheet for a section the user actually enabled -- an empty
"Performance" tab on a report nobody asked to include it would be exactly
the silent padding this product's export code avoids elsewhere. Each sheet
carries only the columns its section has: an Inventory sheet with empty
Severity/Detail columns, or a Security sheet with an empty Cost column, is
padding of the same kind.
"""
from __future__ import annotations

import io
from datetime import datetime
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .csv_renderer import _flatten, coverage_rows, report_facts

_SHEET_TITLES = {
    "executive_summary": "Summary",
    "inventory": "Inventory",
    "security": "Security Findings",
    "utilization": "Utilization",
    "performance": "Performance",
    "recommendations": "Recommendations",
    "cost": "Cost",
    "cost_summary": "Account Spend",
    "anomalies": "Anomalies (Acct-Wide)",
    "certificates": "Expiring Credentials",
    "changes": "Config Changes",
    "posture": "Posture & Compliance",
    "identity": "Identity & Access",
    "activity": "Account Activity",
    "data_quality": "Data Quality",
}

_SUMMARY_SECTIONS = ("executive_summary", "cost_summary", "posture", "identity", "activity",
                     "data_quality")

_ROW_COLUMNS = (
    ("report_section", "Section"),
    ("resource_name", "Resource"), ("provider_resource_id", "Provider Resource ID"),
    ("resource_type", "Type"), ("region", "Region"), ("attribution", "Attribution"),
    ("title", "Title"), ("severity", "Severity"), ("detail", "Detail"),
    ("metric_value", "Metric/Value"), ("status", "Status"),
    ("detected_at", "Detected At"), ("remediation", "Remediation"),
    ("estimated_monthly_savings", "Est. Monthly Savings"), ("savings_currency", "Savings Currency"),
    ("expires_at", "Expires At"), ("days_left", "Days Left"), ("store", "Store"),
    ("in_use", "In Use"),
    ("cost_amount", "Billed Cost"), ("currency", "Currency"),
    ("cost_period_days", "Cost Window (days)"), ("billed_service", "Billed Service"),
    ("cost_fetched_at", "Cost Fetched At"),
    ("environment", "Environment"), ("owner", "Owner"),
    ("resource_group", "Resource Group"), ("compartment", "Compartment"),
    ("tags", "Tags"),
    ("discovered_at", "Discovered At"), ("last_seen_at", "Last Seen At"),
)

_DATE_FIELDS = {"detected_at", "discovered_at", "last_seen_at", "cost_fetched_at", "expires_at"}
_NUMERIC_FIELDS = {"cost_amount", "estimated_monthly_savings", "days_left"}

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill(start_color="1E6FB8", end_color="1E6FB8", fill_type="solid")


def _parse_date(value: Any):
    if not isinstance(value, str):
        return value
    try:
        # Excel's datetime cell type cannot hold a tzinfo -- stripped here,
        # not dropped silently: every timestamp this product stores is UTC.
        return datetime.fromisoformat(value.replace("Z", "+00:00")).replace(tzinfo=None)
    except ValueError:
        return value


def _style_header(ws: Worksheet, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}1"


def _cell_value(key: str, value: Any):
    if key in _DATE_FIELDS:
        return _parse_date(value)
    if key in _NUMERIC_FIELDS and value is not None:
        return float(value)
    if isinstance(value, dict):
        return ", ".join(f"{k}={v}" for k, v in value.items()) or None
    if isinstance(value, list):
        return ", ".join(str(v) for v in value) or None
    return value


def _write_row_sheet(wb: Workbook, title: str, rows: List[Dict[str, Any]]) -> None:
    ws = wb.create_sheet(title[:31])
    present = {k for row in rows for k in row}
    columns = [(k, label) for k, label in _ROW_COLUMNS if k in present]
    if not columns:  # an enabled section with no rows still gets its header
        columns = [(k, label) for k, label in _ROW_COLUMNS
                   if k in ("resource_name", "title", "detail", "status")]
    ws.append([label for _, label in columns])
    for row in rows:
        ws.append([_cell_value(key, row.get(key)) for key, _ in columns])
    _style_header(ws, len(columns))
    for col in range(1, len(columns) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 20


# Lists of records inside a summary section that read as a table, not as
# "frameworks.3.passed = 4" lines: (section, key) -> sheet title.
_TABLES = {
    ("posture", "frameworks"): "Compliance",
    ("executive_summary", "top_risks"): "Top Risks",
    ("executive_summary", "top_cost_resources"): "Top Cost",
    ("identity", "top_identities"): "Top Identities",
    ("activity", "top_actors"): "Top Actors",
}


def _write_table_sheet(wb: Workbook, title: str, rows: List[Dict[str, Any]]) -> None:
    ws = wb.create_sheet(title[:31])
    keys = list(dict.fromkeys(k for row in rows for k in row))
    ws.append([k.replace("_", " ").capitalize() for k in keys])
    for row in rows:
        ws.append([_cell_value(k, row.get(k)) for k in keys])
    _style_header(ws, len(keys))
    for col in range(1, len(keys) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 24


def _write_summary_sheet(wb: Workbook, dataset: Dict[str, Any]) -> None:
    ws = wb.create_sheet("Summary")
    ws.append(["Field", "Value"])
    for key, value in report_facts(dataset):
        ws.append([key, value])
    for section in _SUMMARY_SECTIONS:
        data = dataset.get("sections", {}).get(section)
        if not isinstance(data, dict):
            continue
        ws.append([])
        ws.append([_SHEET_TITLES.get(section, section)])
        flat = {k: v for k, v in data.items() if (section, k) not in _TABLES}
        for key, value in _flatten(flat):
            ws.append([key, str(value) if isinstance(value, (dict, list)) else value])
        for (sec, key), title in _TABLES.items():
            if sec == section and data.get(key):
                ws.append([key, f"See the '{title}' sheet."])
    _style_header(ws, 2)
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 60


def build_xlsx(dataset: Dict[str, Any]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)  # default blank sheet

    _write_summary_sheet(wb, dataset)
    for (section, key), title in _TABLES.items():
        rows = (dataset.get("sections", {}).get(section) or {}).get(key) if isinstance(
            dataset.get("sections", {}).get(section), dict) else None
        if rows:
            _write_table_sheet(wb, title, rows)

    for section, rows in dataset.get("sections", {}).items():
        if section in _SUMMARY_SECTIONS or not isinstance(rows, list):
            continue
        _write_row_sheet(wb, _SHEET_TITLES.get(section, section), rows)

    coverage = coverage_rows(dataset)
    if coverage:
        _write_row_sheet(wb, "Coverage", coverage)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
